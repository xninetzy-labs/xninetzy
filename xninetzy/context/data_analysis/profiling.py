from __future__ import annotations

import csv
import uuid
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from xninetzy.context.data_analysis.dataset import (
    ColumnProfile,
    DatasetSummary,
    checksum_file,
    detect_format,
)


_SEMANTIC_HINTS_NUMERIC = {"int", "float", "decimal"}
_SEMANTIC_HINTS_DATE = {"date", "datetime", "timestamp"}


def _infer_physical_type(values: Iterable[Any]) -> str:
    types = {type(v).__name__ for v in values if v is not None}
    if not types:
        return "empty"
    if types <= {"int", "float"}:
        return "numeric"
    if types <= {"str"}:
        return "string"
    if types <= {"bool"}:
        return "boolean"
    if "datetime" in types or "date" in types:
        return "datetime"
    return "mixed"


def _infer_semantic_type(name: str, physical: str, sample: tuple[Any, ...]) -> str:
    n = name.lower()
    if any(tok in n for tok in ("id", "uuid", "key", "code")):
        return "identifier"
    if any(tok in n for tok in ("date", "time", "ts", "created", "updated")):
        return "temporal"
    if physical == "numeric" and any(
        tok in n for tok in ("amount", "price", "revenue", "total", "cost", "value")
    ):
        return "measure"
    if physical == "numeric":
        return "measure"
    if physical == "boolean":
        return "boolean"
    if physical == "datetime":
        return "temporal"
    return "categorical"


def profile_columns(rows: list[list[Any]]) -> list[ColumnProfile]:
    if not rows:
        return []
    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    profiles: list[ColumnProfile] = []
    for col_idx, header in enumerate(headers):
        col_values: list[Any] = []
        for row in rows[1:]:
            if col_idx < len(row):
                col_values.append(row[col_idx])
        non_null = [v for v in col_values if v is not None and v != ""]
        coerced = [_coerce_scalar(v) for v in non_null]
        physical = _infer_physical_type(coerced)
        semantic = _infer_semantic_type(header, physical, tuple(coerced[:20]))
        profiles.append(
            ColumnProfile(
                name=header,
                physical_type=physical,
                semantic_type=semantic,
                nullable=len(non_null) < len(col_values),
                unique_count=len({repr(v) for v in coerced}),
                null_count=len(col_values) - len(non_null),
                example_values=tuple(str(v) for v in coerced[:3]),
            )
        )
    return profiles


def _coerce_scalar(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (int, float, bool)):
        return v
    s = str(v).strip()
    if s == "":
        return None
    try:
        if s.lstrip("-").isdigit():
            return int(s)
        return float(s)
    except ValueError:
        return s


def profile_csv(path: str | Path, max_rows: int = 100_000) -> DatasetSummary:
    p = Path(path)
    rows: list[list[Any]] = []
    with open(p, "r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        for idx, row in enumerate(reader):
            if idx > max_rows:
                break
            rows.append(row)
    if not rows:
        return DatasetSummary(
            dataset_id=f"ds-{uuid.uuid4().hex[:12]}",
            source_uri=str(p),
            format=detect_format(p),
            row_count=0,
            column_count=0,
            checksum=checksum_file(p),
            columns=(),
        )
    columns = profile_columns(rows)
    return DatasetSummary(
        dataset_id=f"ds-{uuid.uuid4().hex[:12]}",
        source_uri=str(p),
        format=detect_format(p),
        row_count=len(rows) - 1,
        column_count=len(columns),
        checksum=checksum_file(p),
        columns=tuple(columns),
    )


def profile_xlsx(
    path: str | Path,
    sheet_name: str | None = None,
    max_rows: int = 100_000,
) -> DatasetSummary:
    p = Path(path)
    workbook = load_workbook(p, read_only=True, data_only=True)
    try:
        ws = workbook[sheet_name] if sheet_name else workbook.active
        rows: list[list[Any]] = []
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx > max_rows:
                break
            rows.append(list(row))
    finally:
        workbook.close()
    if not rows:
        return DatasetSummary(
            dataset_id=f"ds-{uuid.uuid4().hex[:12]}",
            source_uri=str(p),
            format=detect_format(p),
            row_count=0,
            column_count=0,
            checksum=checksum_file(p),
            columns=(),
        )
    columns = profile_columns(rows)
    return DatasetSummary(
        dataset_id=f"ds-{uuid.uuid4().hex[:12]}",
        source_uri=str(p),
        format=detect_format(p),
        row_count=len(rows) - 1,
        column_count=len(columns),
        checksum=checksum_file(p),
        columns=tuple(columns),
    )


def profile_dataset(path: str | Path, **kwargs: Any) -> DatasetSummary:
    fmt = detect_format(path)
    if fmt.value == "csv":
        return profile_csv(path, **kwargs)
    if fmt.value == "xlsx":
        return profile_xlsx(path, **kwargs)
    raise ValueError(f"unsupported format for profiling: {fmt.value}")
