from __future__ import annotations

import hashlib
import json
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


@dataclass(frozen=True, slots=True)
class TableArtifact:
    artifact_id: str
    name: str
    format: str
    schema: tuple[str, ...]
    row_count: int
    source: str
    transformation: str
    path: str
    checksum: str
    created_at: str = field(default_factory=_utcnow)
    validation_status: str = "pending"

    def to_dict(self) -> dict[str, Any]:
        return {
            "artifact_id": self.artifact_id,
            "name": self.name,
            "format": self.format,
            "schema": list(self.schema),
            "row_count": self.row_count,
            "source": self.source,
            "transformation": self.transformation,
            "path": self.path,
            "checksum": self.checksum,
            "created_at": self.created_at,
            "validation_status": self.validation_status,
        }


def _checksum(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(64 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def generate_xlsx_artifact(
    *,
    path: str | Path,
    name: str,
    rows: Sequence[Mapping[str, Any]],
    source: str = "",
    transformation: str = "",
) -> TableArtifact:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    schema: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in schema:
                schema.append(str(key))
    if not schema:
        schema = ["value"]
    workbook = Workbook()
    try:
        ws = workbook.active
        ws.title = name[:31] or "data"
        ws.append(schema)
        for row in rows:
            ws.append([row.get(col) for col in schema])
        workbook.save(p)
    finally:
        workbook.close()
    return TableArtifact(
        artifact_id=f"art-{uuid.uuid4().hex[:12]}",
        name=name,
        format="xlsx",
        schema=tuple(schema),
        row_count=len(rows),
        source=source,
        transformation=transformation,
        path=str(p),
        checksum=_checksum(p),
    )


def validate_xlsx_artifact(artifact: TableArtifact) -> TableArtifact:
    p = Path(artifact.path)
    if not p.exists():
        return TableArtifact(**{**artifact.to_dict(), "validation_status": "missing"})
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return TableArtifact(**{**artifact.to_dict(), "validation_status": "corrupt"})
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return TableArtifact(**{**artifact.to_dict(), "validation_status": "empty"})
    header = [str(c) for c in rows[0]] if rows[0] else []
    expected = list(artifact.schema)
    if header != expected:
        return TableArtifact(
            **{**artifact.to_dict(), "validation_status": "schema_mismatch"}
        )
    actual_rows = len(rows) - 1
    if actual_rows != artifact.row_count:
        return TableArtifact(
            **{**artifact.to_dict(), "validation_status": "row_count_mismatch"}
        )
    return TableArtifact(**{**artifact.to_dict(), "validation_status": "valid"})


def artifact_to_json(artifact: TableArtifact) -> str:
    return json.dumps(artifact.to_dict(), indent=2, ensure_ascii=False)
